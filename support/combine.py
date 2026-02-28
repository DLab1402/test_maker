import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

INPUT_FILE = "J:/My Drive/Cao_Thang/document/TEST_BANK/NNLT/chapter3/question.xlsx"
OUTPUT_FILE = "J:/My Drive/Cao_Thang/document/TEST_BANK/NNLT/chapter3/question_combined.xlsx"

# ---------- read excel ----------
df = pd.read_excel(INPUT_FILE)

# ---------- combine columns ----------
newline = "\r\n"

df["Question"] = (
    df["Question"].astype(str) + newline +
    "A. " + df["A"].astype(str) + newline +
    "B. " + df["B"].astype(str) + newline +
    "C. " + df["C"].astype(str) + newline +
    "D. " + df["D"].astype(str)
)

# ---------- drop old columns ----------
df.drop(columns=["A", "B", "C", "D"], inplace=True)

# ---------- save ----------
df.to_excel(OUTPUT_FILE, index=False)

# ---------- enable wrap text ----------
wb = load_workbook(OUTPUT_FILE)
ws = wb.active

# find Question column index
question_col = None
for i, cell in enumerate(ws[1], start=1):
    if cell.value == "Question":
        question_col = i
        break

# apply wrap text
if question_col:
    for row in ws.iter_rows(min_col=question_col, max_col=question_col):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)

wb.save(OUTPUT_FILE)

print("Done. File saved as:", OUTPUT_FILE)