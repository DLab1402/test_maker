import os
import re
import webbrowser
from pathlib import Path
from docx import Document
from PyQt6.QtWidgets import QFileDialog


from setting import SOURCE_TYPE, DRIVE_LINK, LOCAL_PATH, get_path

class compose:
    def __init__(self,GUI,dataload):
        self.path = dataload.path
        self.data = dataload.chap_tray
        
        self.GUI = GUI
        self.GUI.preview.clicked.connect(self.make)
        self.GUI.save.clicked.connect(self.save)

    def save(self):
        file_path, _ = QFileDialog.getSaveFileName(
            self.GUI,
            "Save File",          
            "test.docx",      
            "Word Files (*.docx)"
        )
        if file_path != None:
            doc = self.make_word()
            if doc != None:
                doc.save(file_path)

    def make(self):
        html = self.preview_html()
        path = Path(get_path("template\\preview.html"))
        path.write_text(html, encoding="utf-8")

        webbrowser.open(path.resolve().as_uri())

    def count(self):
        TL = []
        TN = []
        for key,chap in self.data.items():
            try:
                print(chap)
                if chap["ques_choose"] == None:
                    continue
                for i in chap["ques_choose"]["TL"]:
                    print(i)
                    TL.append({"chap":key, "Content": chap["chap_content"][i]["Content"], 
                            "Question": chap["chap_content"][i]["Question"],
                            
                            "Link": chap["chap_content"][i]["Link"]}) 
                for i in chap["ques_choose"]["TN"]:
                    print(i)
                    TN.append({"chap":key,"Content": chap["chap_content"][i]["Content"], 
                            "Question": chap["chap_content"][i]["Question"],
                            
                            "Link": chap["chap_content"][i]["Link"]})
            except Exception as e:
                print("Error in compose.count():", e)

        return TN,TL

    def preview_html(self):
        html = """
        <html>
        <body style="
            font-family: 'Times New Roman';
            font-size: 12pt;
            background: white;
            width: 800px;
            margin: 40px auto;
            line-height: 1.6;
        ">
        """
        TN, TL = self.count()
        html += "<h2 style='text-align:center; margin-top:40px;'>PHẦN TRẮC NGHIỆM</h2>"
        for i, q in enumerate(TN, start=1):
            html += f"<p><b>Câu {i}:</b></p>"
            content = re.split(r'\s(?=[A-D]\.)', q['Question'])
            content = "<br>".join(content)
            content = content.replace("_x000D_", "")
            html += f"<p>{content}</p>"
            
            try:
                if isinstance(q['Link'],str):
                    print(q['Link'])
                    link = os.path.join(self.path,self.GUI.subject_list.currentText(),q["chap"],'image',q['Link'])
                    if Path(link).exists():
                        html += f'<img src="{link}" alt="Image" style="max-width:600px;"><br>'
            except MemoryError as e:
                pass
            html += "<br>"
        html += "<h2 style='text-align:center; margin-top:40px;'>PHẦN TỰ LUẬN</h2>"
        for i, q in enumerate(TL, start=1):
            html += f"<p><b>Câu {i}:</b></p>"
            content = re.split(r'\s(?=[A-D]\.)', q['Question'])
            content = "<br>".join(content)
            content = content.replace("_x000D_", "")
            html += f"<p>{content}</p>"
            
            try:
                if isinstance(q['Link'],str):
                    link = os.path.join(self.path,self.GUI.subject_list.currentText(),q["chap"],'image',q['Link'])
                    if Path(link).exists():
                        html += f'<img src="{link}" alt="Image" style="max-width:600px;"><br>'
            except MemoryError as e:
                pass
            html += "<br>"

        html += "</body></html>"
        return html
    
    def make_word(self):
        doc = Document(get_path("template\\test_template.docx"))
        TN, TL = self.count()
        
        doc.add_paragraph("Trắc Nghiệm", style="Heading 1") 
        for i, q in enumerate(TN, start = 1):
            print(q)
            doc.add_paragraph(f"", style="Heading 2")
            content = q["Question"].replace("_x000D_", "")
            doc.add_paragraph(f"{content}", style="Normal")
            try:
                if isinstance(q['Link'],str):
                    print(q['Link'])
                    link = os.path.join(self.path,self.GUI.subject_list.currentText(),q["chap"],'image',q['Link'])
                    if Path(link).exists():
                        doc.add_picture(link)
            except MemoryError as e:
                pass

            doc.add_paragraph("")
        doc.add_paragraph("Tự Luận", style="Heading 1")  
        for i, q in enumerate(TL, start = 1):
            doc.add_paragraph(f"", style="Heading 2")
            content = q["Question"].replace("_x000D_", "")
            doc.add_paragraph(f"{content}", style="Normal")
            try:
                if isinstance(q['Link'],str):
                    print(q['Link'])
                    link = os.path.join(self.path,self.GUI.subject_list.currentText(),q["chap"],'image',q['Link'])
                    if Path(link).exists():
                        doc.add_picture(link)
            except MemoryError as e:
                pass

            doc.add_paragraph("")
        
        return doc